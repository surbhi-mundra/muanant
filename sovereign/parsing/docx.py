"""DOCX parser — structure-preserving using python-docx.

Preserves: headings (with level), paragraphs, tables, metadata.
"""

from __future__ import annotations

from typing import ClassVar

import docx
from docx.document import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from sovereign.parsing.model import (
    Block,
    DocumentMetadata,
    Page,
    ParsedDocument,
    Table,
    TableCell,
    TableRow,
)


class DOCXParser:
    """Parse DOCX files preserving headings, paragraphs, and tables."""

    NAME: ClassVar[str] = "docx.python-docx"
    SUPPORTED: ClassVar[list[str]] = [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ]

    @property
    def name(self) -> str:
        return self.NAME

    def supported_mimes(self) -> list[str]:
        return list(self.SUPPORTED)

    def parse(self, data: bytes, filename: str = "") -> ParsedDocument:
        warnings: list[str] = []
        blocks: list[Block] = []
        metadata = DocumentMetadata()

        import io

        try:
            doc: DocxDocument = docx.Document(io.BytesIO(data))
        except Exception as e:
            warnings.append(f"failed to open DOCX: {e}")
            return ParsedDocument(
                source_filename=filename,
                source_mime_type=self.SUPPORTED[0],
                parser_name=self.NAME,
                parse_warnings=warnings,
            )

        # Metadata
        cp = doc.core_properties
        metadata.title = cp.title or None
        metadata.author = cp.author or None
        metadata.subject = cp.subject or None
        if cp.created:
            metadata.created_at = cp.created.isoformat()
        if cp.modified:
            metadata.modified_at = cp.modified.isoformat()

        # Walk the document body in order (paragraphs + tables interleaved)
        section_path: list[str] = []
        total_words = 0

        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                para = Paragraph(element, doc)
                text = para.text.strip()
                if not text:
                    continue

                style_name = (para.style.name or "").lower() if para.style else ""

                if "heading" in style_name or "title" in style_name:
                    # Extract heading level
                    level = 1
                    if "heading" in style_name:
                        try:
                            level = int(style_name.split("heading")[-1].strip())
                        except ValueError:
                            level = 1
                    elif "title" in style_name:
                        level = 0

                    if level > 0 and level <= len(section_path):
                        section_path = section_path[: level - 1]
                    section_path.append(text)

                    blocks.append(
                        Block(
                            kind="heading",
                            text=text,
                            section_path=list(section_path[:-1]),
                            level=max(level, 1),
                        )
                    )
                elif "list" in style_name:
                    total_words += len(text.split())
                    blocks.append(
                        Block(
                            kind="list_item",
                            text=text,
                            section_path=list(section_path),
                        )
                    )
                else:
                    total_words += len(text.split())
                    blocks.append(
                        Block(
                            kind="paragraph",
                            text=text,
                            section_path=list(section_path),
                        )
                    )

            elif tag == "tbl":
                table = _parse_table(DocxTable(element, doc), section_path)
                blocks.append(
                    Block(
                        kind="table",
                        text=table.markdown or "",
                        section_path=list(section_path),
                        table=table,
                    )
                )
                total_words += sum(len(c.text.split()) for r in table.rows for c in r.cells)

        metadata.word_count = total_words
        metadata.char_count = sum(len(b.text) for b in blocks)
        metadata.page_count = 1  # DOCX doesn't have real pages

        # DOCX has no real pages — wrap in a single synthetic page
        pages = [Page(page_number=1, blocks=blocks, is_synthetic=True)]

        return ParsedDocument(
            source_filename=filename,
            source_mime_type=self.SUPPORTED[0],
            pages=pages,
            metadata=metadata,
            parser_name=self.NAME,
            parse_warnings=warnings,
        )


def _parse_table(docx_table: DocxTable, section_path: list[str]) -> Table:
    """Convert a python-docx Table to our Table model."""
    rows: list[TableRow] = []
    for i, row in enumerate(docx_table.rows):
        cells = []
        for cell in row.cells:
            cells.append(
                TableCell(
                    text=cell.text.strip(),
                    is_header=(i == 0),  # first row as header
                )
            )
        rows.append(TableRow(cells=cells))

    # Build markdown representation
    md_lines: list[str] = []
    if rows:
        # Header
        header_cells = [c.text for c in rows[0].cells]
        md_lines.append("| " + " | ".join(header_cells) + " |")
        md_lines.append("| " + " | ".join("---" for _ in header_cells) + " |")
        for row in rows[1:]:
            md_lines.append("| " + " | ".join(c.text for c in row.cells) + " |")

    return Table(rows=rows, markdown="\n".join(md_lines))

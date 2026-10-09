"""Integration tests for document ingestion end-to-end."""

from __future__ import annotations

import io

import docx
import pytest
from openpyxl import Workbook

from sovereign.ingestion.service import (
    delete_document,
    get_document,
    get_parsed_document,
    ingest_document,
    list_documents,
    validate_document,
)
from sovereign.storage.db.base import init_schema, reset_engine, session_scope
from sovereign.storage.objects import reset_object_store


@pytest.fixture(autouse=True)
def _fresh_env(monkeypatch: pytest.MonkeyPatch, tmp_path: pytest.TempPathFactory):
    """Each test gets a fresh SQLite DB + fresh object store."""
    monkeypatch.setenv("DATABASE_URL", "sqlite:///:memory:")
    monkeypatch.setenv("SOVEREIGN_ENV", "dev")
    monkeypatch.setenv("OBJECT_STORE_FS_ROOT", str(tmp_path / "objects"))
    from sovereign.core.config import reset_settings_cache

    reset_settings_cache()
    reset_engine()
    reset_object_store()
    init_schema()
    yield
    reset_engine()
    reset_object_store()
    reset_settings_cache()


PROJECT_ID = "01JQTESTPROJECT0000000001"


# ---------------------------------------------------------------------------
# Test documents
# ---------------------------------------------------------------------------
def _make_txt() -> bytes:
    return (
        b"Introduction\n\nThis is a test document about pumps.\n"
        b"It has multiple paragraphs.\n\nCONCLUSION\n\nThe pump is operational."
    )


def _make_md() -> bytes:
    return (
        b"# Inspection Report\n\n## Findings\n\nPump P-101 shows wear.\n"
        b"\n## Recommendation\n\nSchedule maintenance."
    )


def _make_csv() -> bytes:
    return b"Equipment,Status,Last_Inspected\nPump-101,OK,2024-01-15\nValve-202,Leak,2024-02-20\n"


def _make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Name"])
    ws.append([1, "Alpha"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Test Doc", 0)
    d.add_paragraph("Content here.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Validation tests
# ---------------------------------------------------------------------------
def test_validate_txt() -> None:
    data = _make_txt()
    result = validate_document(data, "test.txt", "text/plain")
    assert result.is_valid
    assert result.mime_type == "text/plain"
    assert len(result.sha256) == 64
    assert result.size_bytes == len(data)


def test_validate_empty_file() -> None:
    result = validate_document(b"", "empty.txt", "text/plain")
    assert not result.is_valid
    assert "empty" in (result.quarantine_reason or "")


def test_validate_oversized() -> None:
    data = b"x" * (201 * 1024 * 1024)  # 201 MB
    result = validate_document(data, "big.bin", "application/octet-stream", max_size_mb=200)
    assert not result.is_valid
    assert "exceeds" in (result.quarantine_reason or "")


def test_validate_detects_pdf_magic() -> None:
    data = b"%PDF-1.4\n%test pdf content"
    result = validate_document(data, "test.pdf", "application/octet-stream")
    assert result.is_valid
    assert result.mime_type == "application/pdf"


def test_validate_detects_docx_magic() -> None:
    # DOCX is a ZIP — create a real one
    data = _make_docx()
    result = validate_document(data, "test.docx", "application/octet-stream")
    assert result.is_valid
    assert "wordprocessingml" in result.mime_type


def test_validate_detects_xlsx_magic() -> None:
    data = _make_xlsx()
    result = validate_document(data, "test.xlsx", "application/octet-stream")
    assert result.is_valid
    assert "spreadsheetml" in result.mime_type


# ---------------------------------------------------------------------------
# Ingestion tests
# ---------------------------------------------------------------------------
def test_ingest_txt_document() -> None:
    result = ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    assert result.status == "parsed"
    assert result.document_id
    assert result.mime_type == "text/plain"
    assert result.version == 1
    assert not result.is_duplicate


def test_ingest_creates_document_row() -> None:
    result = ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    doc = get_document(PROJECT_ID, result.document_id)
    assert doc.id == result.document_id
    assert doc.original_filename == "test.txt"
    assert doc.status == "parsed"
    assert doc.sha256 == result.sha256


def test_ingest_dedup() -> None:
    """Uploading the same content twice should return the existing doc."""
    data = _make_txt()
    r1 = ingest_document(PROJECT_ID, "test.txt", data, "text/plain")
    r2 = ingest_document(PROJECT_ID, "test.txt", data, "text/plain")
    assert r1.document_id == r2.document_id
    assert r2.is_duplicate
    assert r2.status == "duplicate"


def test_ingest_versioning() -> None:
    """Different content with same filename should increment version."""
    ingest_document(PROJECT_ID, "report.txt", b"version 1 content", "text/plain")
    r2 = ingest_document(PROJECT_ID, "report.txt", b"version 2 different content", "text/plain")
    assert r2.version == 2
    assert r2.document_id  # new doc ID


def test_ingest_md_document() -> None:
    result = ingest_document(PROJECT_ID, "test.md", _make_md(), "text/markdown")
    assert result.status == "parsed"
    assert result.mime_type == "text/markdown"


def test_ingest_csv_document() -> None:
    result = ingest_document(PROJECT_ID, "test.csv", _make_csv(), "text/csv")
    assert result.status == "parsed"


def test_ingest_xlsx_document() -> None:
    result = ingest_document(PROJECT_ID, "test.xlsx", _make_xlsx(), "application/octet-stream")
    assert result.status == "parsed"
    assert "spreadsheetml" in result.mime_type


def test_ingest_docx_document() -> None:
    result = ingest_document(PROJECT_ID, "test.docx", _make_docx(), "application/octet-stream")
    assert result.status == "parsed"
    assert "wordprocessingml" in result.mime_type


def test_ingest_stores_parsed_document() -> None:
    """The parsed document JSON should be retrievable after ingest."""
    result = ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    parsed = get_parsed_document(PROJECT_ID, result.document_id)
    assert parsed.source_filename == "test.txt"
    assert len(parsed.pages) >= 1
    assert parsed.metadata.word_count is not None
    assert parsed.metadata.word_count > 0


def test_ingest_empty_file_quarantines() -> None:
    """Empty files should be quarantined, not stored."""
    result = ingest_document(PROJECT_ID, "empty.txt", b"", "text/plain")
    assert result.status == "quarantined"
    assert result.quarantine_reason
    assert not result.document_id  # no doc created


# ---------------------------------------------------------------------------
# List / delete
# ---------------------------------------------------------------------------
def test_list_documents() -> None:
    ingest_document(PROJECT_ID, "a.txt", b"content a", "text/plain")
    ingest_document(PROJECT_ID, "b.txt", b"content b", "text/plain")
    docs = list_documents(PROJECT_ID)
    assert len(docs) == 2


def test_list_documents_project_isolation() -> None:
    """Documents from project A should not appear in project B's list."""
    ingest_document(PROJECT_ID, "a.txt", b"content a", "text/plain")
    ingest_document("PROJ_OTHER", "b.txt", b"content b", "text/plain")
    docs_a = list_documents(PROJECT_ID)
    docs_b = list_documents("PROJ_OTHER")
    assert len(docs_a) == 1
    assert len(docs_b) == 1
    assert docs_a[0].original_filename == "a.txt"
    assert docs_b[0].original_filename == "b.txt"


def test_delete_document() -> None:
    """Deleting should remove from DB and object store."""
    result = ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    doc_id = result.document_id
    delete_document(PROJECT_ID, doc_id)
    from sovereign.core.errors import NotFoundError

    with pytest.raises(NotFoundError):
        get_document(PROJECT_ID, doc_id)


def test_delete_cross_project_fails() -> None:
    """Deleting a doc from the wrong project should fail."""
    result = ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    from sovereign.core.errors import NotFoundError

    with pytest.raises(NotFoundError):
        delete_document("PROJ_OTHER", result.document_id)


# ---------------------------------------------------------------------------
# Audit trail
# ---------------------------------------------------------------------------
def test_ingest_writes_audit_event() -> None:
    """Each ingest should produce an audit event."""
    from sovereign.audit.log import verify_full_chain

    ingest_document(PROJECT_ID, "test.txt", _make_txt(), "text/plain")
    with session_scope() as s:
        from sqlalchemy import text

        count = s.execute(
            text("SELECT COUNT(*) FROM audit_events WHERE category = 'ingestion'")
        ).scalar()
        assert count is not None and count >= 1
        # Chain should still be valid
        assert verify_full_chain(s) is True
